from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from src.screener.engine import (
    build_screener_dataset,
    latest_company_records,
    load_config,
    run_preset,
)
from src.screener.ranking import build_ranking_scores


OUTPUT_PATH = Path("output/screener_output.xlsx")

RANKING_COLUMNS = [
    "profitability_score",
    "growth_score",
    "valuation_score",
    "composite_score",
    "component_count",
    "score_completeness",
    "ranking_status",
    "overall_rank",
    "sector_rank",
]

EXPORT_COLUMNS = [
    "overall_rank",
    "sector_rank",
    "company_id",
    "company_name",
    "year",
    "broad_sector",
    *RANKING_COLUMNS[:6],
    "ranking_status",
    "engine_return_on_equity_pct",
    "engine_net_profit_margin_pct",
    "engine_operating_profit_margin_pct",
    "engine_roce_pct",
    "engine_roa_pct",
    "engine_debt_to_equity",
    "engine_interest_coverage",
    "engine_asset_turnover",
    "free_cash_flow_cr",
    "cash_from_operations_cr",
    "capex_cr",
    "capex_intensity_pct",
    "fcf_conversion_pct",
    "cfo_pat_ratio",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "earnings_per_share",
    "book_value_per_share",
    "dividend_payout_ratio_pct",
    "pe_ratio",
    "pb_ratio",
    "ev_ebitda",
    "dividend_yield_pct",
    "market_cap_crore",
    "sales",
    "net_profit",
]


def prepare_sheet(df: pd.DataFrame) -> pd.DataFrame:
    columns = list(dict.fromkeys(c for c in EXPORT_COLUMNS if c in df.columns))
    result = df[columns].copy()

    if "composite_score" in result.columns:
        result = result.sort_values(
            ["composite_score", "company_id"],
            ascending=[False, True],
            na_position="last",
        )

    return result.reset_index(drop=True)


def format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[1].height = 32

        headers = {
            cell.value: cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        score_col = headers.get("composite_score")
        if score_col and worksheet.max_row > 1:
            letter = worksheet.cell(1, score_col).column_letter
            worksheet.conditional_formatting.add(
                f"{letter}2:{letter}{worksheet.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )

        for header, column in headers.items():
            letter = worksheet.cell(1, column).column_letter
            if header in {
                "composite_score",
                "profitability_score",
                "growth_score",
                "valuation_score",
                "score_completeness",
                "engine_return_on_equity_pct",
                "engine_net_profit_margin_pct",
                "engine_operating_profit_margin_pct",
                "engine_roce_pct",
                "engine_roa_pct",
                "dividend_yield_pct",
                "revenue_cagr_5yr",
                "pat_cagr_5yr",
                "eps_cagr_5yr",
            }:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row, column).number_format = "0.00"

        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = max(
                (len(str(c.value)) for c in column_cells if c.value is not None),
                default=10,
            )
            worksheet.column_dimensions[letter].width = min(
                max(max_length + 2, 12), 28
            )


def main() -> None:
    print("=" * 65)
    print("NIFTY100 — SCREENER EXCEL EXPORT")
    print("=" * 65)

    config = load_config()
    dataset = build_screener_dataset()
    latest = latest_company_records(dataset)
    ranked = build_ranking_scores(latest)

    sheets: dict[str, pd.DataFrame] = {"Ranking": ranked}

    for preset_name in config["presets"]:
        result = run_preset(preset_name)

        # Avoid duplicate ranking columns before joining current rankings.
        result = result.drop(
            columns=[c for c in RANKING_COLUMNS if c in result.columns],
            errors="ignore",
        )

        available_rank_columns = [
            c for c in RANKING_COLUMNS if c in ranked.columns
        ]
        result = result.merge(
            ranked[["company_id", *available_rank_columns]],
            on="company_id",
            how="left",
            validate="many_to_one",
        )

        sort_columns = [
            c for c in ["composite_score", "company_id"]
            if c in result.columns
        ]
        if sort_columns:
            result = result.sort_values(
                sort_columns,
                ascending=[False if c == "composite_score" else True
                           for c in sort_columns],
                na_position="last",
            )

        # Use the configured preset key as the worksheet name.
        sheets[preset_name[:31]] = result

        print(f"{preset_name}: {len(result)} companies")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            prepare_sheet(frame).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )
        format_workbook(writer)

    print()
    print(f"Created: {OUTPUT_PATH}")
    print(f"Sheets: {len(sheets)}")
    print(f"Ranking rows: {len(ranked)}")
    print("Missing financial values remain blank; they are not replaced with zero.")


if __name__ == "__main__":
    main()
