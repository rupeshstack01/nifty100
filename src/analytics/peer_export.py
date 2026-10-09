from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule


DB_PATH = Path("data/processed/nifty100.db")
OUTPUT_PATH = Path("output/peer_comparison.xlsx")


METRICS = [
    "ROE",
    "ROCE",
    "NPM",
    "D/E",
    "FCF",
    "PAT CAGR 5Y",
    "Revenue CAGR 5Y",
    "EPS CAGR 5Y",
    "Interest Coverage",
    "Asset Turnover",
]


def load_data():
    with sqlite3.connect(DB_PATH) as conn:
        percentiles = pd.read_sql_query(
            "SELECT * FROM peer_percentiles",
            conn,
        )

        companies = pd.read_sql_query(
            "SELECT * FROM company_summary",
            conn,
        )

    return percentiles, companies


def build_metric_table(percentiles: pd.DataFrame) -> pd.DataFrame:
    """
    Convert long peer_percentiles table into one row per company,
    containing the 10 required metric values and percentile ranks.
    """

    value = percentiles.pivot_table(
        index=[
            "company_id",
            "peer_group_name",
            "year",
        ],
        columns="metric",
        values="value",
        aggfunc="first",
    ).reset_index()

    percentile = percentiles.pivot_table(
        index=[
            "company_id",
            "peer_group_name",
            "year",
        ],
        columns="metric",
        values="percentile_rank",
        aggfunc="first",
    ).reset_index()

    value_columns = [
        "company_id",
        "peer_group_name",
        "year",
    ]

    for metric in METRICS:
        if metric in value.columns:
            value_columns.append(metric)

    value = value[value_columns]

    percentile_columns = [
        "company_id",
    ]

    for metric in METRICS:
        if metric in percentile.columns:
            percentile_columns.append(metric)

    percentile = percentile[percentile_columns]

    percentile = percentile.rename(
        columns={
            metric: f"{metric} Percentile"
            for metric in METRICS
            if metric in percentile.columns
        }
    )

    result = value.merge(
        percentile,
        on="company_id",
        how="left",
    )

    return result


def find_company_name_column(
    companies: pd.DataFrame,
) -> str | None:

    for column in [
        "company_name",
        "name",
        "company",
    ]:
        if column in companies.columns:
            return column

    return None


def build_sheet(
    peer_group: str,
    group: pd.DataFrame,
    companies: pd.DataFrame,
) -> pd.DataFrame:

    company_name_column = find_company_name_column(
        companies
    )

    if company_name_column is not None:

        names = companies[
            ["company_id", company_name_column]
        ].drop_duplicates("company_id")

        names = names.rename(
            columns={
                company_name_column: "company_name"
            }
        )

        group = group.merge(
            names,
            on="company_id",
            how="left",
        )

    else:
        group["company_name"] = group["company_id"]

    columns = [
        "company_id",
        "company_name",
    ]

    # 10 metric values.
    for metric in METRICS:
        if metric in group.columns:
            columns.append(metric)

    # 10 percentile values.
    for metric in METRICS:
        percentile_column = f"{metric} Percentile"

        if percentile_column in group.columns:
            columns.append(percentile_column)

    output = group[columns].copy()

    output = output.sort_values(
        "company_id"
    ).reset_index(drop=True)

    # --------------------------------------------------------
    # Median row
    # --------------------------------------------------------

    median_row = {
        "company_id": "MEDIAN",
        "company_name": "Peer Group Median",
    }

    for column in columns[2:]:
        median_row[column] = output[column].median(
            numeric_only=True
        )

    output = pd.concat(
        [
            output,
            pd.DataFrame([median_row]),
        ],
        ignore_index=True,
    )

    # --------------------------------------------------------
    # Benchmark row
    # --------------------------------------------------------
    #
    # The workbook requirement calls for a benchmark row.
    # We use the peer-group median as the benchmark.
    # This is explicitly labelled "BENCHMARK".
    # --------------------------------------------------------

    benchmark_row = {
        "company_id": "BENCHMARK",
        "company_name": f"{peer_group} Benchmark",
    }

    for column in columns[2:]:
        benchmark_row[column] = output[column].median(
            numeric_only=True
        )

    # Put benchmark before median.
    median_index = output.index[
        output["company_id"].eq("MEDIAN")
    ]

    if len(median_index):
        insert_at = int(median_index[0])
    else:
        insert_at = len(output)

    output = pd.concat(
        [
            output.iloc[:insert_at],
            pd.DataFrame([benchmark_row]),
            output.iloc[insert_at:],
        ],
        ignore_index=True,
    )

    return output


def format_workbook(path: Path) -> None:

    wb = load_workbook(path)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    benchmark_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD966",
    )

    median_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3",
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFEB9C",
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E1F2",
        )
    )

    for ws in wb.worksheets:

        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        for cell in ws[1]:

            cell.fill = header_fill
            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        ws.row_dimensions[1].height = 30

        # ----------------------------------------------------
        # Rows
        # ----------------------------------------------------

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
        ):

            company_id = row[0].value

            if company_id == "BENCHMARK":

                for cell in row:
                    cell.fill = benchmark_fill
                    cell.font = Font(
                        bold=True
                    )

            elif company_id == "MEDIAN":

                for cell in row:
                    cell.fill = median_fill
                    cell.font = Font(
                        bold=True
                    )

            for cell in row:

                cell.border = thin_border

                cell.alignment = Alignment(
                    vertical="center"
                )

        # ----------------------------------------------------
        # Percentile conditional formatting
        # ----------------------------------------------------

        for column_index, header in enumerate(
            ws[1],
            start=1,
        ):

            if "Percentile" not in str(
                header.value
            ):
                continue

            column_letter = header.column_letter

            data_range = (
                f"{column_letter}2:"
                f"{column_letter}{ws.max_row}"
            )

            # >= 75 = green
            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="greaterThanOrEqual",
                    formula=["75"],
                    fill=green_fill,
                ),
            )

            # 25 <= percentile < 75 = yellow
            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="between",
                    formula=["25", "74.999999"],
                    fill=yellow_fill,
                ),
            )

            # <= 25 = red
            ws.conditional_formatting.add(
                data_range,
                CellIsRule(
                    operator="lessThanOrEqual",
                    formula=["25"],
                    fill=red_fill,
                ),
            )

        # ----------------------------------------------------
        # Number formats
        # ----------------------------------------------------

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
        ):

            for cell in row:

                if isinstance(
                    cell.value,
                    (int, float),
                ):

                    header = str(
                        ws.cell(
                            1,
                            cell.column,
                        ).value
                    )

                    if "Percentile" in header:
                        cell.number_format = "0.0"

                    else:
                        cell.number_format = "0.00"

        # ----------------------------------------------------
        # Column widths
        # ----------------------------------------------------

        for column_cells in ws.columns:

            column_letter = (
                column_cells[0].column_letter
            )

            max_length = 0

            for cell in column_cells:

                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            ws.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                28,
            )

    wb.save(path)


def main() -> None:

    print("=" * 70)
    print("D20 — PEER COMPARISON EXCEL EXPORT")
    print("=" * 70)

    percentiles, companies = load_data()

    print(
        f"Peer percentile rows: "
        f"{len(percentiles)}"
    )

    print(
        f"Companies: "
        f"{percentiles['company_id'].nunique()}"
    )

    metric_table = build_metric_table(
        percentiles
    )

    named_groups = sorted(
        metric_table.loc[
            metric_table["peer_group_name"]
            != "No peer group assigned",
            "peer_group_name",
        ]
        .dropna()
        .unique()
    )

    print(
        f"Named peer groups: "
        f"{len(named_groups)}"
    )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:

        for peer_group in named_groups:

            group = metric_table[
                metric_table["peer_group_name"]
                .eq(peer_group)
            ].copy()

            sheet = build_sheet(
                peer_group,
                group,
                companies,
            )

            sheet.to_excel(
                writer,
                sheet_name=peer_group[:31],
                index=False,
            )

            print(
                f"  {peer_group}: "
                f"{len(group)} companies"
            )

    format_workbook(
        OUTPUT_PATH
    )

    # --------------------------------------------------------
    # Validation
    # --------------------------------------------------------

    wb = load_workbook(
        OUTPUT_PATH,
        read_only=True,
    )

    print()
    print("Validation")
    print("-" * 70)

    print(
        "Sheets:",
        len(wb.sheetnames),
    )

    total_failures = 0

    if len(wb.sheetnames) != 11:
        print(
            "FAIL: Expected 11 sheets"
        )
        total_failures += 1

    expected_metric_columns = 22

    for sheet_name in wb.sheetnames:

        ws = wb[sheet_name]

        print(
            f"{sheet_name}: "
            f"{ws.max_row - 1} data rows, "
            f"{ws.max_column} columns"
        )

        if ws.max_column != expected_metric_columns:
            print(
                f"  FAIL: expected "
                f"{expected_metric_columns} columns"
            )
            total_failures += 1

        headers = [
            cell.value
            for cell in ws[1]
        ]

        if "company_id" not in headers:
            total_failures += 1

        if "company_name" not in headers:
            total_failures += 1

        percentile_count = sum(
            "Percentile" in str(header)
            for header in headers
        )

        if percentile_count != 10:
            print(
                f"  FAIL: expected 10 percentile columns, "
                f"found {percentile_count}"
            )
            total_failures += 1

        company_count = (
            ws.max_row - 1 - 2
        )

        if company_count < 1:
            print(
                "  FAIL: no company records"
            )
            total_failures += 1

    wb.close()

    if total_failures:
        print(
            f"\nD20 STATUS: FAIL "
            f"({total_failures} validation failures)"
        )

    else:
        print("\nD20 STATUS: PASS")

    print(
        f"Output: {OUTPUT_PATH}"
    )


if __name__ == "__main__":
    main()